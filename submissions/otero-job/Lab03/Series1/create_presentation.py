"""
create_presentation.py
======================
Generates the Series1_Results_Presentation.pptx PowerPoint presentation
using python-pptx.

The presentation explains the results of the three exercises:
    1. Convergence of (1 + 1/n)^n toward e
    2. Convergence of (a^h - 1) / h toward ln(a)
    3. Approximation of e^x using the power series

Output: Series1_Results_Presentation.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ----------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PPTX = os.path.join(SCRIPT_DIR, "Series1_Results_Presentation.pptx")

# Image paths
IMG1 = os.path.join(SCRIPT_DIR, "exercise1_results.png")
IMG2 = os.path.join(SCRIPT_DIR, "exercise2_results.png")
IMG3 = os.path.join(SCRIPT_DIR, "exercise3_results.png")

# Colors
DARK_BLUE = RGBColor(0x1A, 0x52, 0x76)
MEDIUM_BLUE = RGBColor(0x29, 0x80, 0xB9)
LIGHT_BLUE = RGBColor(0xEB, 0xF5, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x2C, 0x3E, 0x50)
LIGHT_GRAY = RGBColor(0x7F, 0x8C, 0x8D)
ACCENT_GREEN = RGBColor(0x1E, 0x84, 0x49)
ACCENT_PURPLE = RGBColor(0x6C, 0x34, 0x83)

# Fonts
FONT_TITLE = "Calibri"
FONT_BODY = "Calibri"

# ----------------------------------------------------------------------
# Helper functions
# ----------------------------------------------------------------------
def add_title_slide(prs):
    """Slide 1: Title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(9.0), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Series1 Numerical Methods Laboratory"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9.0), Inches(1.0))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Numerical Series and Convergence Analysis"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
    p2.font.name = FONT_BODY
    p2.alignment = PP_ALIGN.CENTER

    # Placeholder text
    txBox3 = slide.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(9.0), Inches(1.5))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Course: Numerical Methods Laboratory"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
    p3.font.name = FONT_BODY
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf3.add_paragraph()
    p4.text = "Student: [Student Name]"
    p4.font.size = Pt(16)
    p4.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
    p4.font.name = FONT_BODY
    p4.alignment = PP_ALIGN.CENTER

    p5 = tf3.add_paragraph()
    p5.text = "Date: [Date]"
    p5.font.size = Pt(16)
    p5.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
    p5.font.name = FONT_BODY
    p5.alignment = PP_ALIGN.CENTER


def add_content_slide(prs, title, bullets, title_color=DARK_BLUE):
    """Generic content slide with title and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # Title bar
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = title_color
    p.font.name = FONT_TITLE

    # Horizontal line
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.1), Inches(9.0), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = MEDIUM_BLUE
    line.line.fill.background()

    # Bullet points
    txBox2 = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(8.6), Inches(5.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(12)
        p.level = 0

    return slide


def add_image_slide(prs, title, img_path, bullets, title_color=DARK_BLUE):
    """Slide with title, image, and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # Title bar
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = title_color
    p.font.name = FONT_TITLE

    # Horizontal line
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.1), Inches(9.0), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = MEDIUM_BLUE
    line.line.fill.background()

    # Image on the left
    if os.path.exists(img_path):
        pic = slide.shapes.add_picture(img_path, Inches(0.5), Inches(1.5), width=Inches(5.5))
    else:
        print(f"WARNING: Image not found: {img_path}")

    # Bullet points on the right
    txBox2 = slide.shapes.add_textbox(Inches(6.3), Inches(1.5), Inches(3.5), Inches(5.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = bullet
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(8)
        p.level = 0

    return slide


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------
def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Slide 1: Title ---
    add_title_slide(prs)

    # --- Slide 2: Introduction ---
    add_content_slide(
        prs,
        "Introduction",
        [
            "Numerical approximation is the process of finding approximate solutions to mathematical problems using computational methods.",
            "Convergence refers to the tendency of a sequence or series to approach a specific value as the number of terms or iterations increases.",
            "Mathematical series are fundamental in numerical methods because they allow us to compute transcendental functions (e, ln, e^x) to arbitrary precision.",
            "This laboratory investigates three key convergence phenomena:",
            "   • (1 + 1/n)^n → e",
            "   • (a^h - 1)/h → ln(a)",
            "   • e^x = Σ(x^n / n!)",
        ],
    )

    # --- Slide 3: Exercise 1 Theory ---
    add_content_slide(
        prs,
        "Exercise 1: Theory — (1 + 1/n)^n → e",
        [
            "The mathematical constant e ≈ 2.718281828459045 is one of the most important constants in mathematics.",
            "The sequence a_n = (1 + 1/n)^n converges to e as n approaches infinity.",
            "For small n, the value is far from e:",
            "   • n = 1: (1 + 1/1)^1 = 2.0",
            "   • n = 2: (1 + 1/2)^2 = 2.25",
            "   • n = 4: (1 + 1/4)^4 = 2.4414",
            "As n increases, the value approaches e more closely.",
            "The error |(1 + 1/n)^n - e| decreases approximately as 1/n.",
        ],
    )

    # --- Slide 4: Exercise 1 Results ---
    add_image_slide(
        prs,
        "Exercise 1: Results",
        IMG1,
        [
            "The bar chart shows (1 + 1/n)^n approaching the dashed line at e.",
            "At n = 1,048,576, the value is 2.718280532282396.",
            "The absolute error at this n is 1.30 × 10⁻⁶.",
            "The error decreases approximately linearly on a log-log scale.",
            "This demonstrates steady, predictable convergence toward e.",
        ],
    )

    # --- Slide 5: Exercise 2 Theory ---
    add_content_slide(
        prs,
        "Exercise 2: Theory — (a^h - 1)/h → ln(a)",
        [
            "The natural logarithm ln(a) is the inverse of the exponential function e^x.",
            "The difference quotient (a^h - 1)/h approximates the derivative of a^x at x = 0.",
            "As h approaches zero, this quotient converges to ln(a).",
            "This is a direct application of the definition of the derivative:",
            "   f'(0) = lim_{h→0} (a^h - a^0)/h = lim_{h→0} (a^h - 1)/h = ln(a)",
            "We test this for a = 2, a = e, and a = 3.",
            "A convergence tolerance of 1 × 10⁻⁶ is used.",
        ],
    )

    # --- Slide 6: Exercise 2 Results ---
    add_image_slide(
        prs,
        "Exercise 2: Results",
        IMG2,
        [
            "The grouped bar chart shows (a^h - 1)/h for a = 2, e, 3.",
            "As h decreases from 0.1 to 10⁻⁶, each approximation approaches its respective ln(a).",
            "All three values reach the 1 × 10⁻⁶ tolerance at h = 10⁻⁶.",
            "The error decreases approximately linearly with h.",
            "This confirms first-order convergence of the difference quotient.",
        ],
    )

    # --- Slide 7: Exercise 3 Theory ---
    add_content_slide(
        prs,
        "Exercise 3: Theory — e^x = Σ(x^n / n!)",
        [
            "The exponential function e^x can be expressed as an infinite power series:",
            "   e^x = Σ_{n=0}^{∞} (x^n / n!)",
            "A partial sum S_N = Σ_{n=0}^{N} (x^n / n!) approximates e^x.",
            "As N increases, S_N converges to e^x.",
            "To avoid numerical overflow from huge factorials, we use the recurrence:",
            "   term_{n+1} = term_n × x / (n + 1)",
            "   partial_sum = partial_sum + term",
            "This is numerically stable and supports N up to 10,000.",
        ],
    )

    # --- Slide 8: Exercise 3 Results ---
    add_image_slide(
        prs,
        "Exercise 3: Results",
        IMG3,
        [
            "The bar chart shows partial sums S_N approaching e² = 7.389056098930650.",
            "Error < 10⁻² is achieved at N = 7 terms.",
            "Error < 10⁻⁶ is achieved at N = 13 terms.",
            "Error < 10⁻¹⁰ is achieved at N = 17 terms.",
            "The final error at N = 10,000 is 1.78 × 10⁻¹⁵ (machine precision).",
        ],
    )

    # --- Slide 9: Numerical Accuracy and Stability ---
    add_content_slide(
        prs,
        "Numerical Accuracy and Stability",
        [
            "Floating-point precision limits the accuracy of numerical computations.",
            "Directly computing huge factorials (e.g., 10000!) causes overflow.",
            "The recurrence relation term_{n+1} = term_n × x / (n + 1) avoids this problem.",
            "Convergence tolerance (1 × 10⁻⁶) provides a practical stopping criterion.",
            "Numerical stability ensures that small errors do not grow during computation.",
            "These principles are essential for reliable scientific computing.",
        ],
    )

    # --- Slide 10: Conclusions ---
    add_content_slide(
        prs,
        "Conclusions",
        [
            "Exercise 1: (1 + 1/n)^n converges to e as n increases.",
            "Exercise 2: (a^h - 1)/h converges to ln(a) as h approaches zero.",
            "Exercise 3: The power series Σ(x^n / n!) converges to e^x as N increases.",
            "Numerical methods can approximate mathematical functions accurately.",
            "Increasing resolution (n or h) or the number of terms (N) generally improves approximation accuracy.",
            "Numerical stability and floating-point precision are critical considerations.",
        ],
    )

    # Save the presentation
    prs.save(OUTPUT_PPTX)
    print(f"Presentation saved as: {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()